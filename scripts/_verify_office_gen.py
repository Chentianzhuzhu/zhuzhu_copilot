# -*- coding: utf-8 -*-
"""office 三件套高质量生成器 回归冒烟（真实样例，无 mock）。

覆盖：
1. create 三件：Word(封面/目录/表格/标题层级/引用)、PPT(封面/两栏/卡片/图表/图形/强调页)、
   Excel(多表/合计/原生图表/数字格式)
2. execute_tool 接线：以 create_pptx 走一遍 agent 工具分发路径（模拟 LLM 调用参数）
3. 文本回读：extract_text 能读回中文内容
4. beautify 三件：对副本各美化一次确保不抛异常

产物输出到 output/office_showcase/（供人工打开验收视觉质量）。
"""
import os
import shutil
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

OUT = ROOT / "output" / "office_showcase"
OUT.mkdir(parents=True, exist_ok=True)

FAIL = []


def check(name, cond, extra=""):
    print(("PASS  " if cond else "FAIL  ") + name + (f"  {extra}" if extra else ""))
    if not cond:
        FAIL.append(name)


def main():
    from winapp_migrator.office import (beautify_docx, beautify_pptx,
                                        beautify_xlsx, build_docx,
                                        build_pptx, build_xlsx)
    from winapp_migrator.core import agent_tools

    # ---------- 1. Word：项目方案（tech 主题，封面/目录/表格/引用） ----------
    docx_path = OUT / "智慧校园解决方案.docx"
    r = build_docx(
        str(docx_path),
        title="智慧校园一体化解决方案",
        style={
            "theme": "tech",
            "cover": {"subtitle": "连接校园每一处 · 让管理与教学更高效",
                      "author": "晨天智联科技", "date": "2026 年 9 月",
                      "image": ""},
            "header_text": "晨天智联 · 智慧校园解决方案",
        },
        paragraphs=[
            "[toc]",
            "# 一、项目背景",
            "随着校园信息化建设的深入，业务系统林立、数据孤岛、管理分散等问题日益突出。",
            "本项目以统一平台 + 智能终端的方式，重构校园数字化底座：",
            "- 统一身份认证与数据中台，消除重复建设",
            "- 以 AI 助手承载教务、后勤、安防等高频事项",
            "> 目标：让 80% 的日常事务性工作通过自助化、智能化方式完成，释放一线人员精力。",
            "## 1.1 现状与痛点",
            "1. 系统分散：全校现有业务系统 12 套，账号口令各自独立",
            "2. 数据不通：成绩、课表、消费数据难以联动分析",
            "3. 服务被动：大量人工咨询与线下跑腿",
            "# 二、总体方案",
            "| 模块 | 面向对象 | 核心能力 | 交付形态 |",
            "| 智慧门户 | 全校师生 | 统一入口、应用聚合 | Web + App |",
            "| AI 助手 | 师生 | 智能问答、事务代办、主动提醒 | 大模型应用 |",
            "| 数据中台 | 管理部门 | 数据治理、可视化分析 | 平台服务 |",
            "| 物联感知 | 后勤安保 | 门禁、能耗、环境监测 | 智能终端 |",
            "## 2.1 实施节奏",
            "- 第一阶段：门户与数据中台（第 1-3 个月）",
            "- 第二阶段：AI 助手与高频应用（第 4-6 个月）",
            "# 三、预期成效",
            "通过统一平台建设，预计可显著降低系统运维成本、提升师生办事效率，形成可复制的",
            "智慧校园建设样板。[pagebreak]",
            "# 四、结语",
            "我们期待与校方共同打造安全、高效、有温度的智慧校园。",
        ],
    )
    check("build_docx 返回", isinstance(r, dict) and "已生成" in r.get("text", ""))

    # ---------- 2. PPT：产品发布（black-gold，两栏/卡片/图表/图形/强调页） ----------
    pptx_path = OUT / "晨天助手产品发布.pptx"
    slides = [
        {"title": "为什么要做晨天助手",
         "bg_color": "FAFAF7",
         "bullets": ["校园服务仍停留在排队、窗口、电话的传统模式，体验与效率双低。",
                     "师生高频诉求高度集中：查课表、办证明、报修、请假、缴费。",
                     "AI 大模型成熟，让自然语言办事第一次成为可能。",
                     "我们决定做一个真正'办事'的校园助手，而不是聊天玩具。"]},
        {"title": "产品架构：三层一体",
         "layout": "two_col",
         "columns": [
             {"title": "能力层", "bullets": ["统一身份与数据服务", "事务引擎：接单、分派、闭环", "大模型问答与知识库"]},
             {"title": "交互层", "bullets": ["App / 公众号 / 大屏多端一致", "语音、文字、拍照多模态", "主动提醒与代办中心"]}]},
        {"title": "四大核心能力",
         "anim": {"graphics": "circle"},
         "cards": [
             {"title": "智慧问答", "desc": "政策制度、办事指南即问即答，24 小时在线"},
             {"title": "一键办事", "desc": "请销假、证明开具、报修缴费全流程在线闭环"},
             {"title": "主动服务", "desc": "根据课表与日程主动提醒，重要节点不错过"},
             {"title": "数据看板", "desc": "管理者实时掌握运行数据，决策有据可依"}]},
        {"title": "与常见方案的对比",
         "table": {"title": "能力对比一览",
                   "header": ["能力维度", "传统门户", "普通问答机器人", "晨天助手"],
                   "rows": [["事务闭环", "不支持", "不支持", "端到端支持"],
                            ["主动服务", "无", "无", "日程驱动提醒"],
                            ["多模态交互", "网页表单", "文本问答", "语音/文字/拍照"],
                            ["数据沉淀", "独立", "无", "统一看板"]]}},
        {"title": "上线三个月，体验指标持续向好",
         "chart": {"type": "line", "title": "服务响应时长（分钟）",
                   "labels": ["第 1 周", "第 2 周", "第 3 周", "第 4 周", "第 6 周", "第 8 周", "第 12 周"],
                   "values": [28, 21, 17, 12, 8, 5, 3]}},
        {"title": "应用场景：从入学到毕业",
         "diagram": {"type": "flow", "title": "校园全生命周期服务",
                     "items": ["迎新报到", "学习生活", "考试实习", "毕业离校"]}},
        {"title": "我们要成为的那件事",
         "layout": "center_highlight",
         "highlight": "不是替代师生，而是替他们跑腿、算数、记日程",
         "bullets": ["把重复劳动交给系统，把时间还给教学与成长",
                     "从一所学校出发，做成可复制的智慧校园样板"]},
        {"title": "产品节奏与邀请",
         "anim": {"exit": "fade"},
         "chart": {"type": "pie", "title": "首批种子校试用计划",
                   "labels": ["需求共建", "试点运行", "正式发布"],
                   "values": [3, 2, 1]},
         "bullets": ["诚邀 3 所伙伴学校深度共建",
                     "欢迎现场体验演示环境"]},
    ]
    r = build_pptx(str(pptx_path), title="晨天助手", slides=slides,
                   style={"theme": "black-gold",
                          "cover": {"subtitle": "让校园服务像聊天一样简单",
                                    "author": "晨天智联", "date": "2026 · 秋季发布会"},
                          "footer": "晨天智联 · 内部材料",
                          "transition": "random"})
    check("build_pptx 返回", isinstance(r, dict) and "已生成" in r.get("text", ""))

    # ---------- 3. Excel：经营数据分析（多表 + 合计 + 原生图表） ----------
    xlsx_path = OUT / "2026上半年经营分析.xlsx"
    r = build_xlsx(
        str(xlsx_path),
        sheets=[
            {"name": "订单明细",
             "wordart": {"text": "2026 上半年订单明细（示例）"},
             "rows": [
                 ["订单号", "区域", "渠道", "金额", "毛利率"],
                 ["SO-26001", "华东", "直营", {"v": 128000, "format": "#,##0"}, {"v": 0.32, "format": "0.0%"}],
                 ["SO-26002", "华北", "分销", {"v": 96500, "format": "#,##0"}, {"v": 0.27, "format": "0.0%"}],
                 ["SO-26003", "华南", "电商", {"v": 143200, "format": "#,##0"}, {"v": 0.35, "format": "0.0%"}],
                 ["SO-26004", "西南", "直营", {"v": 76400, "format": "#,##0"}, {"v": 0.29, "format": "0.0%"}],
                 ["SO-26005", "华中", "分销", {"v": 88700, "format": "#,##0"}, {"v": 0.25, "format": "0.0%"}],
                 ["SO-26006", "华东", "电商", {"v": 112500, "format": "#,##0"}, {"v": 0.33, "format": "0.0%"}],
                 ["合计", "", "", "~sum", ""],
             ]},
            {"name": "区域汇总",
             "wordart": {"text": "区域销售汇总"},
             "rows": [
                 ["区域", "销售额", "占比"],
                 ["华东", 240500, 0.33],
                 ["华北", 96500, 0.13],
                 ["华南", 143200, 0.19],
                 ["西南", 76400, 0.10],
                 ["华中", 88700, 0.12],
                 ["其他", 95000, 0.13],
             ],
             "charts": [{"type": "bar", "title": "区域销售额对比",
                         "labels": ["华东", "华北", "华南", "西南", "华中", "其他"],
                         "values": [240500, 96500, 143200, 76400, 88700, 95000]}]},
            {"name": "月度趋势",
             "rows": [["月份", "销售额"],
                      ["1 月", 152000], ["2 月", 128000], ["3 月", 176000],
                      ["4 月", 169000], ["5 月", 198000], ["6 月", 221000]],
             "charts": [{"type": "line", "title": "月度销售额趋势",
                         "labels": ["1 月", "2 月", "3 月", "4 月", "5 月", "6 月"],
                         "values": [152000, 128000, 176000, 169000, 198000, 221000]}]},
        ],
        style={"theme": "business"})
    check("build_xlsx 返回", isinstance(r, dict) and "已生成" in r.get("text", ""))

    # ---------- 4. execute_tool 接线冒烟（走 agent 参数分发路径） ----------
    handler_pptx = OUT / "_handler_pptx.pptx"
    rr = agent_tools.execute_tool("create_pptx", {
        "path": str(handler_pptx), "title": "接线验证",
        "slides": [{"title": "接线正常", "bullets": ["create_pptx 经 execute_tool 调用成功"]}],
        "style": {"theme": "business"}})
    check("execute_tool 接线", isinstance(rr, dict) and "已生成" in rr.get("text", ""), rr.get("text", "")[:60])

    # ---------- 5. 回读校验 ----------
    for p in (docx_path, pptx_path, xlsx_path):
        check(f"产物存在且为 zip: {p.name}", p.is_file() and zipfile.is_zipfile(str(p)))
    for p, keys in ((docx_path, ["智慧校园", "目录", "总体方案", "晨天智联"]),
                    (pptx_path, ["晨天助手", "对比", "架构", "晨天智联"]),
                    (xlsx_path, ["订单明细", "区域汇总", "月度趋势"])):
        if p.suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(p))
            missing = [k for k in keys if k not in wb.sheetnames]
            check(f"xlsx 多工作表完整 {p.name}", not missing,
                  ("缺:" + ",".join(missing)) if missing else f"共 {len(wb.sheetnames)} 表")
            continue
        txt = ""
        try:
            rr = agent_tools._extract_text(str(p))
            txt = (rr or {}).get("text") or ""
        except Exception as e:
            check(f"extract_text 回读 {p.name}", False, str(e))
            continue
        missing = [k for k in keys if k not in txt]
        check(f"extract 内容含关键信息 {p.name}", not missing, ("缺:" + ",".join(missing)) if missing else f"共 {len(txt)} 字")

    # ---------- 5.5 PPT 动画/切换结构断言 ----------
    with zipfile.ZipFile(str(pptx_path)) as z:
        import re as _re
        slide_names = sorted(
            (n for n in z.namelist() if _re.match(r"ppt/slides/slide\d+\.xml$", n)),
            key=lambda x: int(_re.search(r"(\d+)", x).group(1)))
        timing_cnt = 0
        trans_cnt = 0
        exit_on_last = ""
        for i, n in enumerate(slide_names):
            xml = z.read(n).decode("utf-8")
            if "<p:timing>" in xml:
                timing_cnt += 1
            if "<p:transition" in xml:
                trans_cnt += 1
            if i == len(slide_names) - 1 and 'transition="out"' in xml:
                exit_on_last = xml.count('p:animEffect')
        check("每页含元素动画 timing", timing_cnt == len(slide_names),
              f"{timing_cnt}/{len(slide_names)}")
        check("每页含切换 transition", trans_cnt == len(slide_names),
              f"{trans_cnt}/{len(slide_names)}")
        check("末页含退场动画", exit_on_last != "", "animEffect 退场链已注入")

    # ---------- 6. beautify 副本冒烟 ----------
    tmp = OUT / "_beautify_tmp"
    tmp.mkdir(exist_ok=True)
    pairs = [
        (docx_path, "beautify_docx", beautify_docx, {"theme": "tech"}),
        (pptx_path, "beautify_pptx", beautify_pptx, {"theme": "black-gold"}),
        (xlsx_path, "beautify_xlsx", beautify_xlsx, {"theme": "business"}),
    ]
    for src, nm, fn, style in pairs:
        cp = tmp / ("cp_" + src.name)
        shutil.copyfile(src, cp)
        try:
            rr = fn(str(cp), style)
            check(f"{nm} 冒烟", isinstance(rr, dict) and "已美化" in rr.get("text", ""), rr.get("text", "")[:50])
        except Exception as e:
            check(f"{nm} 冒烟", False, repr(e))
    shutil.rmtree(tmp, ignore_errors=True)
    try:
        (OUT / "_handler_pptx.pptx").unlink(missing_ok=True)
    except Exception:
        pass

    print("\n==== 结果 ====")
    if FAIL:
        print("FAILED:", len(FAIL))
        for f in FAIL:
            print(" -", f)
        raise SystemExit(1)
    print("全部通过；产物目录:", OUT)


if __name__ == "__main__":
    main()
