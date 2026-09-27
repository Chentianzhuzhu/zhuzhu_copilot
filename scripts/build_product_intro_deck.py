# -*- coding: utf-8 -*-
"""重建「智能产品发布会」10 页 deck（修正版）并渲染为自包含 HTML 演示文稿。

背景：原 deck（产品介绍PPT_v2.pptx）的保真预览 HTML 存在下列缺陷，本脚本按
「内容逐字取自原稿、只修缺陷」的原则重建：

1. 第 5 页（四、市场表现）表格与柱状图数据打架：
   表格「用户数(万)」12/28/45（2022-2024），柱状图 120/280/450/680（2022-2025）。
   → 以表格为准确认口径，柱状图改为 12/28/45（2022-2024），标题改「用户数（万）」。
2. 末页标题 / 艺术字为 #FFFFFF，落在纯白底上不可见 → 改用主色 #0066CC。
3. Markdown 残留：'### 行业地位'、'#### 电商零售'、'**智能识别**' 等原样排进页面
   → 统一写成语义化层级（'## ' 子标题 / '- ' 二级要点 / '**粗体**'）。
4. 全幅 .sg 文本框压住卡片/表格/柱状图 → 生成器排版已修（卡片高度自适应 + 要点收口）。

用法：python scripts/build_product_intro_deck.py [--src 原 pptx] [--out-deck 输出 pptx]
"""

import argparse
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "src"))

from winapp_migrator.office import pptx_builder, preview  # noqa: E402

DECK_TITLE = "智能产品发布会"
DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
DEFAULT_SRC = os.path.join(DESKTOP, "产品介绍PPT_v2.pptx")
DEFAULT_DECK = os.path.join(DESKTOP, "产品介绍PPT_v3.pptx")
DEFAULT_HTML = os.path.join(DESKTOP, "产品介绍PPT_v3.html")

LIGHT_THEME = {"bg": "#E9EDF4", "card": "#FFFFFF", "text": "#1E293B",
               "dim": "#64748B", "border": "#CBD5E1", "accent": "#1F3A5F"}

STYLE = {
    "font_name": "微软雅黑",
    "title_size": 26,
    "body_size": 18,
    "title_color": "0066CC",
    "text_color": "2B3A4A",
    "accent": "0E5A8A",
    "transition": "fade",
    # 封面白底 + 居中标题（与原稿一致；标题字号/颜色可配，不写死在生成器里）
    "cover_style": "centered",
    "cover": {"title_size": 40, "title_color": "0E5A8A"},
}

# 内容取自原 deck（产品介绍PPT_v2.pptx）逐页文案，仅做缺陷修正
SLIDES = [
    {   # 原 slide 2
        "title": DECK_TITLE,
        "bullets": ["全新智能解决方案", "重新定义用户体验", "引领未来科技趋势"],
        "image": "__PIC__",
        "wordart": {"text": "智能未来 触手可及", "size": 44, "color": "0066CC"},
    },
    {   # 原 slide 3 —— 一、产品概述
        "title": "一、产品概述",
        "cards": [
            {"title": "智能引擎", "desc": "搭载最新AI算法，实时学习用户习惯"},
            {"title": "极速响应", "desc": "毫秒级响应速度，流畅无延迟"},
            {"title": "安全可靠", "desc": "企业级数据加密，隐私全方位保护"},
        ],
        "bullets": [
            "## 核心定位",
            "新一代智能产品，融合AI技术与人性化设计，为用户带来前所未有的使用体验",
            "## 目标用户",
            "面向科技爱好者、企业用户及追求高效生活的现代人群",
            "## 核心价值",
            "效率提升50%以上，操作简化60%，学习成本降低80%",
        ],
    },
    {   # 原 slide 4 —— 二、核心功能
        "title": "二、核心功能",
        "diagram": {"type": "flow", "title": "核心技术流程",
                    "items": ["数据采集", "智能分析", "结果呈现", "持续优化"]},
        "bullets": [
            "**智能识别**：AI图像识别准确率高达99.2%，支持100+品类识别。应用场景：商品分类、质量检测、安全监控等",
            "**语音交互**：多轮对话能力，支持方言识别与语义理解。应用场景：智能客服、语音助手、语音控制等",
            "**数据分析**：多维度数据可视化，自动生成分析报告。应用场景：商业洞察、运营决策、用户画像等",
        ],
    },
    {   # 原 slide 5 —— 三、技术优势
        "title": "三、技术优势",
        "cards": [
            {"title": "自主研发", "desc": "核心技术自主可控"},
            {"title": "多端协同", "desc": "手机/PC/云端无缝同步"},
            {"title": "开放生态", "desc": "开放API，支持二次开发"},
        ],
        "bullets": [
            "自研深度学习框架，训练效率提升3倍。底层架构完全自主可控，无第三方依赖",
            "支持私有化部署与云端服务双重模式，满足不同客户的灵活需求",
            "多模态融合技术，视觉、语音、文本统一处理，打破信息孤岛，实现全场景覆盖",
        ],
    },
    {   # 原 slide 6 —— 四、市场表现（表格与柱状图口径已对齐）
        "title": "四、市场表现",
        "table": {"header": ["指标", "2022", "2023", "2024"],
                  "rows": [["营收(亿)", 1.2, 2.8, 4.5],
                           ["用户数(万)", 12, 28, 45],
                           ["NPS", 72, 81, 87]]},
        "chart": {"type": "column", "title": "用户数（万）",
                  "labels": ["2022", "2023", "2024"], "values": [12, 28, 45]},
        "bullets": [
            "## 行业地位",
            "连续3年保持细分市场占有率第一，累计服务客户超过10万家",
            "## 客户评价",
            "NPS净推荐值高达87分，远超行业平均水平",
            "## 增长趋势",
            "年营收增长率保持在45%以上，客户复购率92%",
        ],
    },
    {   # 原 slide 7 —— 五、应用场景
        "title": "五、应用场景",
        "cards": [
            {"title": "智慧零售", "desc": "智能导购+库存管理"},
            {"title": "在线教育", "desc": "个性化学习路径规划"},
            {"title": "智慧城市", "desc": "城市管理智能化升级"},
            {"title": "工业互联网", "desc": "生产效率智能监控"},
        ],
        "bullets": [
            "## 电商零售",
            "智能导购、库存管理、个性化推荐，助力商家提升转化率30%",
            "## 在线教育",
            "个性化学习路径规划，智能辅导提升学习效果",
            "## 智慧城市",
            "城市管理智能化升级，公共服务效率倍增",
        ],
    },
    {   # 原 slide 8 —— 六、发展规划
        "title": "六、发展规划",
        "diagram": {"type": "cycle", "title": "持续发展循环",
                    "items": ["技术创新", "产品迭代", "市场拓展", "用户服务"]},
        "bullets": [
            "## 2025年目标",
            "- 用户规模突破100万",
            "- 进入3个新细分市场",
            "- 推出企业级解决方案",
            "## 2026年愿景",
            "- 构建完整智能生态体系",
            "- 国际化布局，进军东南亚市场",
            "- 发布下一代产品，实现技术迭代",
        ],
    },
    {   # 原 slide 9 —— 七、合作共赢
        "title": "七、合作共赢",
        "bullets": [
            "## 合作伙伴计划",
            "提供全面技术支持与市场资源，共建智能生态",
            "## 开发者平台",
            "开放API接口，丰富应用场景，激发创新活力",
            "## 培训计划",
            "定期举办技术研讨会，赋能合作伙伴成长",
        ],
        "wordart": {"text": "携手共赢 共创未来", "size": 36, "color": "0066CC"},
    },
    {   # 原 slide 10 —— 谢谢观看（标题/艺术字改为可见的主色）
        "title": "谢谢观看",
        "bullets": ["期待与您合作", "共筑智能新时代"],
        "wordart": {"text": "THANK YOU", "size": 48, "color": "0066CC"},
    },
]


def extract_cover_pic(src: str) -> str:
    """从原 deck 中取出内嵌配图（封面图），落盘为 png 供重建复用。"""
    from pptx import Presentation
    prs = Presentation(src)
    out_dir = os.path.join(tempfile.gettempdir(), "zhuzhu_deck_build")
    os.makedirs(out_dir, exist_ok=True)
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) is not None and "PICTURE" in str(shape.shape_type):
                path = os.path.join(out_dir, f"cover_{shape.shape_id}.png")
                with open(path, "wb") as f:
                    f.write(shape.image.blob)
                return path
    return ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=DEFAULT_SRC, help="原 deck（取封面配图 + 校对内容）")
    ap.add_argument("--out-deck", default=DEFAULT_DECK)
    ap.add_argument("--out-html", default=DEFAULT_HTML)
    ap.add_argument("--workdir", default="")
    args = ap.parse_args()

    pic = extract_cover_pic(args.src) if os.path.isfile(args.src) else ""
    slides = []
    for s in SLIDES:
        item = {k: v for k, v in s.items()}
        if item.get("image") == "__PIC__":
            if pic:
                item["image"] = {"path": pic, "width": 3.48}
            else:
                item.pop("image")
        slides.append(item)

    res = pptx_builder.build_pptx(args.out_deck, DECK_TITLE, slides, STYLE,
                                 workdir=args.workdir)
    print(res["text"])
    html = preview.render_office_html(args.out_deck, theme=LIGHT_THEME)
    with open(args.out_html, "w", encoding="utf-8") as f:
        f.write(html or "")
    print(f"[html] {args.out_html} {len(html or '')} chars")


if __name__ == "__main__":
    main()
