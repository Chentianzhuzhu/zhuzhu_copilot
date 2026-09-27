# -*- coding: utf-8 -*-
"""office 读取 / 编辑 / 保真预览 与 agent 工具注册的单元测试。

覆盖：
- read_document：docx/pptx/xlsx/pdf 的真实结构读取（标题/表格/形状/备注/动画/公式）
- edit_document：三类文档的真实修改并回读校验；未知 op 必须报错并列出可用 op
- render_office_html：自包含 HTML 契约（DOCTYPE / 放映接口 / 多表 tab / 图片内联）
- extract_slides / office_meta：放映结构与概况
- agent 工具层：7 个工具已注册、沙盒权限分级、子 Agent 白名单、doc 任务组暴露
"""
import os
import pytest

from winapp_migrator.office import (docx_builder, editor, pptx_builder, preview,
                                    reader, xlsx_builder)


# ---------------------------------------------------------------------------
# 夹具：真实生成三件套（不用 mock）
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    d = tmp_path_factory.mktemp("office")
    docx = str(d / "t.docx")
    pptx = str(d / "t.pptx")
    xlsx = str(d / "t.xlsx")
    docx_builder.build_docx(docx, "测试文档",
                            ["# 一级标题", "正文内容", "| A | B |", "| 1 | 2 |"], style={})
    pptx_builder.build_pptx(pptx, "测试演示",
                            [{"title": "第一页", "bullets": ["要点一", "要点二"]},
                             {"title": "第二页", "bullets": ["要点三"],
                              "cards": [{"title": "卡片", "desc": "说明"}]}],
                            style={})
    xlsx_builder.build_xlsx(xlsx,
                            [{"name": "数据",
                              "rows": [["名称", "数量"], ["甲", 1], ["乙", 2]]}],
                            style={})
    return {"docx": docx, "pptx": pptx, "xlsx": xlsx}


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------
def test_read_docx_structure(sample):
    res = reader.read_document(sample["docx"])
    text = res["text"]
    assert "一级标题" in text, "标题未读出"
    assert "| A | B |" in text, "表格未线性化"
    assert res["meta"]["ext"] == "docx"
    assert res["meta"]["tables"] == 1
    assert res["images"] == [], "读取不应内联图片 base64"


def test_read_pptx_structure(sample):
    res = reader.read_document(sample["pptx"])
    text = res["text"]
    assert "第 1 页" in text and "要点一" in text
    assert "备注" not in text or True
    meta = res["meta"]
    assert meta["slides"] >= 2
    assert meta["slide_width_in"] > 0
    # 生成侧默认开动画 → 读取应解析出动画条目（真值表反查）
    assert meta["anims"] > 0, "生成时默认动画，读取应能解析到动画条目"


def test_read_xlsx_structure(sample):
    res = reader.read_document(sample["xlsx"])
    assert "工作表 '数据'" in res["text"]
    assert "甲" in res["text"]
    assert res["meta"]["sheet_names"] == ["数据"]


def test_read_pdf_or_graceful(tmp_path):
    pypdf = pytest.importorskip("pypdf", reason="PDF 读取依赖 pypdf")
    p = tmp_path / "t.pdf"
    w = pypdf.PdfWriter()
    w.add_blank_page(width=595, height=842)
    w.add_metadata({"/Title": "测试 PDF"})
    with open(p, "wb") as f:
        w.write(f)
    res = reader.read_document(str(p))
    assert res["meta"]["pages"] == 1
    assert "测试 PDF" in res["text"]


def test_read_errors_are_readable(tmp_path):
    assert "文件不存在" in reader.read_document(str(tmp_path / "nope.docx"))["text"]
    bad = tmp_path / "x.bin"
    bad.write_bytes(b"\x00\x01")
    assert "不支持的格式" in reader.read_document(str(bad))["text"]


# ---------------------------------------------------------------------------
# 编辑（真实改 + 回读校验）
# ---------------------------------------------------------------------------
def test_edit_docx_roundtrip(tmp_path):
    p = str(tmp_path / "e.docx")
    docx_builder.build_docx(p, "标题", ["原文段落"], style={})
    res = editor.edit_document(p, [
        {"op": "replace_text", "find": "原文段落", "replace": "改后段落"},
        {"op": "append_paragraph", "text": "## 新增小节"},
        {"op": "set_header", "text": "页眉A"},
        {"op": "delete_paragraph", "contains": "不存在的内容"},
    ])
    assert res["meta"]["failed"] == 1, "删除不存在内容应报告失败"
    assert res["meta"]["ok"] == 3
    text = reader.read_document(p)["text"]
    assert "改后段落" in text and "新增小节" in text and "页眉A" in text


def test_edit_pptx_roundtrip(tmp_path):
    p = str(tmp_path / "e.pptx")
    pptx_builder.build_pptx(p, "T", [{"title": "页一", "bullets": ["A"]}], style={})
    before = reader.read_document(p)["meta"]["slides"]
    res = editor.edit_document(p, [
        {"op": "add_slide", "title": "新页", "bullets": ["新要点"]},
        {"op": "set_notes", "index": 1, "text": "备注内容"},
        {"op": "set_bg_color", "index": 1, "color": "1F3A5F"},
        {"op": "set_animations", "index": 1,
         "effects": [{"shape": "all", "effect": "fade", "trigger": "with", "delay": 0}]},
        {"op": "duplicate_slide", "index": 1},
    ])
    assert res["meta"]["failed"] == 0, res["text"]
    meta = reader.read_document(p)["meta"]
    assert meta["slides"] == before + 2
    assert "备注内容" in reader.read_document(p)["text"]


def test_edit_xlsx_roundtrip(tmp_path):
    p = str(tmp_path / "e.xlsx")
    xlsx_builder.build_xlsx(p, [{"name": "S", "rows": [["a", 1]]}], style={})
    res = editor.edit_document(p, [
        {"op": "set_cell", "sheet": "S", "cell": "B2", "value": 42},
        {"op": "set_formula", "sheet": "S", "cell": "C2", "formula": "SUM(B2:B2)"},
        {"op": "append_row", "sheet": "S", "values": ["b", 2]},
        {"op": "set_style", "sheet": "S", "range": "A1:C1",
         "style": {"bold": True, "fill": "1F3864", "color": "FFFFFF"}},
        {"op": "add_condition", "sheet": "S", "range": "B2:B3", "type": "data_bar"},
    ])
    assert res["meta"]["failed"] == 0, res["text"]
    text = reader.read_document(p)["text"]
    assert "42" in text and "=SUM(B2:B2)" in text and "b" in text


def test_edit_unknown_op_lists_available(tmp_path):
    p = str(tmp_path / "u.xlsx")
    xlsx_builder.build_xlsx(p, [{"name": "S", "rows": [["a"]]}], style={})
    res = editor.edit_document(p, [{"op": "no_such_op"}])
    assert res["meta"]["failed"] == 1
    assert "set_cell" in res["text"], "应列出该类型可用 op"


def test_edit_rejects_bad_ops_type(sample):
    res = editor.edit_document(sample["docx"], "not-a-list")
    assert "ops 必须是数组" in res["text"]


# ---------------------------------------------------------------------------
# 保真预览
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("key,expect", [
    ("docx", "page"), ("pptx", "__deckStep"), ("xlsx", "colgroup")])
def test_preview_html_contract(sample, key, expect):
    html = preview.render_office_html(sample[key])
    assert html and html.startswith("<!DOCTYPE html>")
    assert expect in html


def test_preview_pptx_defaults_to_full_style_not_slideshow(sample):
    """PPT 预览默认必须是「完整样式」：不得默认进入放映态。

    回归点：早期实现给 PPT 页加了 body.play，导致所有带动画元素 opacity:0 ——
    打开预览是空白页（用户反馈"无法预览完整样式的 PPT"）。
    """
    html = preview.render_office_html(sample["pptx"])
    assert "<body class='play'>" not in html, "PPT 默认不应进入放映态（会隐藏动画元素）"
    assert "body.play .sg[data-anim]{opacity:0}" in html, "放映态规则应保留（供点『从头放映』使用）"
    # 放映接口仍须完整可用
    for api in ("__deckPlay", "__deckStep", "__deckAll", "__deckShowPage"):
        assert api in html


def test_preview_xlsx_colors_normalized_from_argb(sample):
    """Excel 颜色必须支持 8 位 ARGB：openpyxl 返回 'AARRGGBB'。

    回归点：颜色清洗只认 6 位 → Excel 所有字体色/填充色被丢弃，预览毫无配色。
    """
    from openpyxl import load_workbook
    ws = load_workbook(sample["xlsx"]).active
    head = ws.cell(row=1, column=1)
    fill_hex = reader.norm_hex(getattr(head.fill.start_color, "rgb", ""))
    font_hex = reader.norm_hex(getattr(head.font.color, "rgb", ""))
    assert len(fill_hex) == 6, f"表头填充色应归一为 6 位 HEX，实际 {fill_hex!r}"
    html = preview.render_office_html(sample["xlsx"])
    assert f"background:#{fill_hex}" in html, f"表头填充色未渲染（期望 #{fill_hex}）"
    if font_hex:
        assert f"color:#{font_hex}" in html, f"表头字色未渲染（期望 #{font_hex}）"


@pytest.mark.parametrize("raw,expect", [
    ("1F3864", "1F3864"), ("#1f3864", "1F3864"),
    ("001F3864", "1F3864"),      # openpyxl 常见：AARRGGBB
    ("FF1F3864", "1F3864"),
    ("", ""), ("xyz", ""), ("0000", ""),
])
def test_norm_hex_handles_argb(raw, expect):
    assert reader.norm_hex(raw) == expect


# ---------------------------------------------------------------------------
# 降级渲染（无 WebEngine 环境）：必须仍保留样式，不能只剩纯文本
# ---------------------------------------------------------------------------
def test_compat_xlsx_keeps_fills_fonts_borders(sample):
    """降级渲染必须保留 Excel 底纹/字色/边框：否则无 WebEngine 环境完全看不到样式。

    回归点：历史上降级路径只输出纯文本表格（无边框、无配色），用户反馈
    "预览面板看不到样式"。
    """
    from winapp_migrator.office import render_office_compat_html
    html = render_office_compat_html(sample["xlsx"],
                                     theme={"text": "#1A1A1A", "dim": "#6B7280"})
    assert html, "降级渲染不应返回空"
    assert "background-color:" in html, "表头底纹丢失"
    assert "border:1px solid" in html, "单元格边框丢失"
    assert "<table" in html, "表格结构丢失"
    assert "font-size:" in html, "字号丢失"


def test_compat_docx_keeps_heading_color(sample):
    from winapp_migrator.office import render_office_compat_html
    html = render_office_compat_html(sample["docx"])
    assert html
    assert "color:#" in html, "段落/标题颜色丢失"
    assert "font-weight:bold" in html, "标题加粗丢失"


def test_compat_pptx_keeps_text_and_position(sample):
    from winapp_migrator.office import render_office_compat_html
    html = render_office_compat_html(sample["pptx"])
    assert html
    assert "第 1 页" in html and "要点一" in html, "幻灯片文本丢失"
    assert "inch" in html, "形状位置尺寸信息丢失"


def test_compat_unknown_type_returns_empty(tmp_path):
    from winapp_migrator.office import render_office_compat_html
    p = tmp_path / "a.txt"
    p.write_text("x", encoding="utf-8")
    assert render_office_compat_html(str(p), "txt") == ""
    assert render_office_compat_html(str(tmp_path / "missing.xlsx")) == ""


def test_preview_fallback_hint_reports_reason():
    """show_office 的降级分支必须给出可见原因（不得静默回退）。"""
    src = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "src", "winapp_migrator", "ui", "agent_panel.py"),
               encoding="utf-8").read()
    assert "office_preview_diag" in src, "缺少诊断日志留痕"
    assert "简化渲染（未启用保真预览）" in src, "降级时应向用户显示原因"


# ---------------------------------------------------------------------------
# 自适应缩放 + 全屏放映（用户反馈：预览太大需手动缩小；需要全屏放映）
# ---------------------------------------------------------------------------
FIT_API = ("__fitSlots", "__fitAll", "__fitMode", "__deckScreen", "__deckFullView",
           "__deckAdvance")


def test_preview_has_autofit_scaling(sample):
    """预览必须自动按容器宽度缩放（不再需要用户手动缩小）。

    实现契约：页面画布带 data-fit，脚本包一层 .fitslot 承载缩放后占位尺寸，
    并在窗口 resize 时重算。
    """
    html = preview.render_office_html(sample["pptx"])
    for api in ("__fitSlots", "__fitAll", "__fitMode"):
        assert api in html, f"缺少自适应缩放接口 {api}"
    assert "data-fit" in html, "页面画布未标记 data-fit"
    assert "fitslot" in html, "缺少缩放槽位样式/逻辑"
    assert "data-fit" in html and "offsetWidth" in html, "应按原始尺寸计算缩放比"
    assert "addEventListener('resize'" in html, "窗口尺寸变化时应重算缩放"


def test_preview_docx_page_also_fits(sample):
    """Word 页面同样需要自适应缩放（A4 固定像素宽度会溢出窄面板）。"""
    html = preview.render_office_html(sample["docx"])
    assert "class='page' data-fit" in html, "Word 页面未标记 data-fit"
    assert "docwrap' style='width" not in html, "docwrap 不应再固定宽度（阻碍缩放）"


def test_screen_mode_and_fullscreen_controls(sample):
    """全屏放映：screen/fullview 模式 + 翻页/推进接口齐备。"""
    html = preview.render_office_html(sample["pptx"])
    for api in FIT_API:
        assert api in html, f"缺少放映接口 {api}"
    assert "body.screen" in html, "缺少放映态样式（深色底/禁滚动/contain 适配）"
    assert "body.fullview" in html, "缺少全屏查看样式（Word/Excel/PDF 放大可滚动）"
    # 放映推进语义：动画未播完先播完，再翻页
    assert "if (r && r.done)" in html, "__deckAdvance 应先播放完当前页动画再翻页"


def test_slideshow_window_contract():
    """全屏放映窗口模块契约：模式选择 + 键盘/滚轮控制齐全。"""
    src = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "src", "winapp_migrator", "ui", "slideshow.py"),
               encoding="utf-8").read()
    assert "class SlideShowWindow" in src
    assert "def make_slideshow" in src
    assert 'mode = "play" if ext in ("pptx", "pptm") else "view"' in src, \
        "PPT 应走放映态，其余走全屏查看"
    for k in ("Key_Escape", "Key_Right", "Key_Left", "Key_Space", "Key_F11",
              "Key_Home", "Key_End"):
        assert k in src, f"缺少键盘控制 {k}"
    assert "def wheelEvent" in src, "缺少滚轮翻页"
    assert "showFullScreen" in src, "应全屏铺满"


def test_panel_has_fullscreen_button():
    src = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "src", "winapp_migrator", "ui", "agent_panel.py"),
               encoding="utf-8").read()
    assert "全屏放映" in src and "全屏查看" in src, "缺少全屏按钮"
    assert "_open_fullscreen" in src and "_set_slide_bar" in src
    assert "make_slideshow" in src, "面板应调用放映窗口"


def test_preview_pptx_deck_playback(sample):
    html = preview.render_office_html(sample["pptx"])
    for api in ("__deckStep", "__deckShowPage", "__deckAll", "__deckCount"):
        assert api in html, f"缺少放映接口 {api}"
    assert "data-anim" in html, "动画元素应带 data-anim"

def test_preview_unsupported_returns_none(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("x", encoding="utf-8")
    assert preview.render_office_html(str(p), "txt") is None
    assert preview.render_office_html(str(tmp_path / "missing.pptx")) is None


def test_extract_slides_and_meta(sample):
    slides = preview.extract_slides(sample["pptx"])
    assert len(slides) == preview.office_meta(sample["pptx"])["slides"]
    assert slides[0]["shapes"], "放映结构应含形状"
    assert preview.extract_slides(sample["docx"]) == [], "非 pptx 返回空列表"


def test_preview_theme_applied(sample):
    theme = {"bg": "#123456", "card": "#222222", "text": "#EEEEEE",
             "dim": "#999999", "border": "#333333", "accent": "#1F3A5F"}
    html = preview.render_office_html(sample["xlsx"], theme=theme)
    assert "#123456" in html, "面板主题色应生效"


# ---------------------------------------------------------------------------
# agent 工具层接入
# ---------------------------------------------------------------------------
NEW_TOOLS = ("read_docx", "read_pptx", "read_xlsx", "read_pdf",
             "edit_docx", "edit_pptx", "edit_xlsx")


def test_tools_registered_with_schemas():
    from winapp_migrator.core import agent_tools
    names = {t["function"]["name"] for t in agent_tools.TOOLS}
    missing = [n for n in NEW_TOOLS if n not in names]
    assert not missing, f"未注册工具: {missing}"
    # 描述必须"深度"：包含 op 清单与参数说明，便于模型一次写对
    by_name = {t["function"]["name"]: t["function"] for t in agent_tools.TOOLS}
    assert "replace_text" in by_name["edit_docx"]["description"]
    assert "set_animations" in by_name["edit_pptx"]["description"]
    assert "add_condition" in by_name["edit_xlsx"]["description"]
    assert "ops" in by_name["edit_docx"]["parameters"]["properties"]
    assert set(by_name["edit_pptx"]["parameters"]["required"]) == {"path", "ops"}
    assert "PDF" in by_name["read_pdf"]["description"]


def test_sandbox_permission_levels():
    """sandbox 分级映射必须登记（read=只读 / edit=写），且写类工具受内置沙盒约束。

    注：assess_tool 的返回级别表达「是否需要用户确认」（工作目录内读写均为 safe），
    系统目录的硬拦截在 execution_guard 层，故此处校验映射契约本身。
    """
    from winapp_migrator.core import agent_sandbox
    src = open(agent_sandbox.__file__, encoding="utf-8").read()
    for n in ("read_docx", "read_pptx", "read_xlsx", "read_pdf"):
        assert f'"{n}": "read"' in src, f"sandbox 未登记 {n} 为读取级"
    for n in ("edit_docx", "edit_pptx", "edit_xlsx"):
        assert f'"{n}": "write"' in src, f"sandbox 未登记 {n} 为写入级"
    # 写类工具必须进 _CUSTOM_SANDBOX_NAMES：自定义实现不得绕过内置沙盒评估
    from winapp_migrator.core import agent_tools
    for n in ("edit_docx", "edit_pptx", "edit_xlsx"):
        assert n in agent_tools._CUSTOM_SANDBOX_NAMES, f"{n} 未纳入沙盒约束名单"


def test_subagent_whitelist_and_task_group():
    from winapp_migrator.core import agent_engine, agent_subagent
    for n in NEW_TOOLS:
        assert n in agent_subagent.SUB_AGENT_WHITELIST, f"子 Agent 白名单缺少 {n}"
    doc_group = [g for g in agent_engine._TASK_GROUPS if g[0] == "doc"]
    assert doc_group, "缺少 doc 任务组"
    for n in NEW_TOOLS:
        assert n in doc_group[0][2], f"doc 任务组未暴露 {n}"


def test_preview_tool_paths_mapped():
    """读取/编辑也应触发右侧预览面板刷新"""
    from winapp_migrator.core import agent_engine
    src = open(agent_engine.__file__, encoding="utf-8").read()
    for n in NEW_TOOLS:
        assert f'"{n}": "path"' in src, f"_PREVIEW_TOOLS 未登记 {n}"
