# -*- coding: utf-8 -*-
"""PPT 动画回归测试（不依赖 PowerPoint，仅 python-pptx + lxml 读 XML）。

覆盖上一轮修复的关键点：
1. 退场动画可生成：slide.anim.exit / style.exit_effect 都会产生 presetClass="exit" 节点；
2. 动画顺序：标题最先；卡片文字晚于其底板；图表数值标签晚于其柱体（不得先于柱体）；
3. 新增平滑效果 float_left/float_right 可解析（pid=30 官方预设），未知效果名回退 fade 不报错；
4. 生成的 .pptx 可被 python-pptx 正常读回（结构调整未破坏文件结构）。
"""

import os
import sys

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

import pytest

from zhuzhu_Copilot.office.pptx_builder import (
    _ANIM, _anim_preset, _page_steps, _shape_bbox, _shape_text, build_pptx,
)
from pptx import Presentation


def _txt_map(slide):
    return {sh.shape_id: _shape_text(sh).strip() for sh in slide.shapes}


def _title_spid(slide):
    for sh in slide.shapes:
        x, y, w, h = _shape_bbox(sh)
        if _shape_text(sh).strip() and w >= 3.0 and y <= 1.6:
            return sh.shape_id
    return None


def _container_spid(slide, spid):
    """返回包含 spid 中心的非文字大底板（若存在）。"""
    cx = cy = 0.0
    for sh in slide.shapes:
        if sh.shape_id == spid:
            x, y, w, h = _shape_bbox(sh)
            cx, cy = x + w / 2, y + h / 2
            break
    for sh in slide.shapes:
        if _shape_text(sh).strip():
            continue
        x, y, w, h = _shape_bbox(sh)
        if x - 0.06 <= cx <= x + w + 0.06 and y - 0.06 <= cy <= y + h + 0.06:
            return sh.shape_id
    return None


def _exit_count(slide):
    from pptx.oxml.ns import qn
    timing = slide._element.find(qn("p:timing"))
    if timing is None:
        return 0
    return sum(1 for ct in timing.iter(qn("p:cTn"))
               if ct.get("presetClass") == "exit")


def _entr_count(slide):
    from pptx.oxml.ns import qn
    timing = slide._element.find(qn("p:timing"))
    if timing is None:
        return 0
    return sum(1 for ct in timing.iter(qn("p:cTn"))
               if ct.get("presetClass") == "entr")


def test_title_first_and_card_text_after_base(tmp_path):
    out = str(tmp_path / "a.pptx")
    build_pptx(out, slides=[
        {"title": "三大升级",
         "cards": [{"title": "更快", "desc": "响应时间降低 50%"},
                   {"title": "更稳", "desc": "可用性 99.99%"}]},
    ])
    prs = Presentation(out)
    slide = prs.slides[0]
    steps = _page_steps(slide, 0, True, {}, None)
    assert steps, "应有动画步骤"
    # 标题最先
    assert steps[0]["spid"] == _title_spid(slide)
    order = [st["spid"] for st in steps]
    # 卡内文字必须晚于其底板
    for sh in slide.shapes:
        t = _shape_text(sh).strip()
        if t and sh.shape_id != steps[0]["spid"]:
            base = _container_spid(slide, sh.shape_id)
            if base is not None:
                assert order.index(base) < order.index(sh.shape_id), \
                    f"文字 {sh.shape_id} 应晚于其底板 {base}"


def test_chart_value_label_after_its_bar(tmp_path):
    out = str(tmp_path / "b.pptx")
    build_pptx(out, slides=[
        {"title": "增长趋势",
         "chart": {"type": "column", "labels": ["Q1", "Q2", "Q3"],
                   "values": [120, 180, 240]}},
    ])
    prs = Presentation(out)
    slide = prs.slides[0]
    order = [st["spid"] for st in _page_steps(slide, 0, True, {}, None)]
    txt = _txt_map(slide)
    for sp in order:
        if txt.get(sp, "") and txt[sp].replace(".", "").isdigit():
            # 柱体 = 与数值标签水平居中对齐、且位于标签下方的无字图形
            # （标签居中于柱体正上方，比"同 x 起点"更贴合当前排版）
            bx, by, bw, bh = _shape_bbox(next(s for s in slide.shapes
                                              if s.shape_id == sp))
            lcx = bx + bw / 2
            bars = [s.shape_id for s in slide.shapes
                    if not _shape_text(s).strip()
                    and abs((_shape_bbox(s)[0] + _shape_bbox(s)[2] / 2) - lcx) < 0.02
                    and _shape_bbox(s)[1] >= by - 0.02
                    and _shape_bbox(s)[3] > 0.05]
            assert bars and order.index(bars[0]) < order.index(sp), \
                f"数值标签 {sp} 不应先于其柱体 {bars}"


def test_exit_anim_from_slide_anim(tmp_path):
    out = str(tmp_path / "c.pptx")
    build_pptx(out, slides=[
        {"title": "本章小结", "bullets": ["统一平台，一套数据"],
         "anim": {"text": "wipe", "exit": "fade"}},
    ])
    prs = Presentation(out)
    slide = prs.slides[0]
    assert _exit_count(slide) > 0, "slide.anim.exit 应生成退场动画"
    assert _entr_count(slide) > 0


def test_exit_anim_from_style_exit_effect(tmp_path):
    out = str(tmp_path / "d.pptx")
    build_pptx(out, style={"exit_effect": "fly_bottom"}, slides=[
        {"title": "增长趋势",
         "chart": {"type": "column", "labels": ["Q1", "Q2"], "values": [120, 180]}},
    ])
    prs = Presentation(out)
    assert _exit_count(prs.slides[0]) > 0, "style.exit_effect 应生成每页整体退场"


def test_exit_anim_disabled_with_anim_false(tmp_path):
    out = str(tmp_path / "e.pptx")
    build_pptx(out, style={"exit_effect": "fade"}, slides=[
        {"title": "目录", "bullets": ["一、回顾"], "anim": False},
        {"title": "正文页", "bullets": ["内容"]},
    ])
    prs = Presentation(out)
    assert _exit_count(prs.slides[0]) == 0, "anim=False 的页应彻底关闭动画（含退场）"
    assert _exit_count(prs.slides[1]) > 0


def test_new_smooth_float_effects_resolve():
    """float_left/float_right 走官方 pid=30（Float 预设）且入场/退场均可用"""
    for name in ("float_left", "float_right"):
        meta = _anim_preset(name)
        assert meta["presetID"] == 30 and meta["presetClass"] == "entr"
        assert meta["anims"], f"{name} 应有属性动画配方"
        out = _anim_preset(name, exit_mode=True)
        assert out["presetClass"] == "exit" and out["mode"] == "out"
    # 中文别名
    assert _anim_preset("浮入自左")["presetID"] == 30
    assert _anim_preset("平滑浮入")["presetID"] == 30
    # 轮换池包含新增平滑效果
    assert "float_left" in _ANIM


def test_unknown_effect_falls_back_to_fade():
    """未知效果名不报错，一律回退 fade（保持向后兼容）"""
    assert _anim_preset("bogus_effect")["presetID"] == 10
    assert _anim_preset(None) is None
    assert _anim_preset("off") is None


def test_build_pptx_returns_path_and_readable(tmp_path):
    out = str(tmp_path / "f.pptx")
    res = build_pptx(out, title="报告", slides=[
        {"title": "页一", "bullets": ["要点一", "要点二"]},
        {"title": "页二", "cards": [{"title": "A", "desc": "a"}]},
    ], style={"theme": "business"})
    assert "已生成" in res["text"] and out in res["text"]
    prs = Presentation(out)
    assert len(prs.slides) == 3   # 封面 + 2 内容页


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))