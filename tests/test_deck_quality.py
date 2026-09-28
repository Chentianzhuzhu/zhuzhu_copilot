# -*- coding: utf-8 -*-
"""deck 质量回归测试：markdown 残留 / 排版遮挡 / 不可见文字 / 页规格 / 缩放脚本。

对应本轮修复的真实缺陷（防回归）：
1. 要点里的 '### 行业地位'、'#### 电商零售'、'**智能识别**' 原样排进成品
   → 生成器须归一为层级标题 / 行内加粗，成品文案不得残留标记；
2. 卡片区固定高度 + 要点固定起点导致正文块互相压叠 → 须自动收口不重叠；
3. 末页标题/艺术字用 #FFFFFF 落在白底上不可见 → 须自动换可读色；
4. 表格与柱状图同页时要能纵向堆叠（旧实现只渲染其中一个）；
5. HTML 页规格：10 页 1280×720、data-fit、.sg/.tb、页码块 96.0×38.4px、
   slidenum 'n / 10' 连续、缩放脚本完整且只保留当前页占位（避免下方大片空白）。
"""
import os
import re
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

import pytest
from pptx import Presentation
from pptx.util import Emu

from zhuzhu_Copilot.office import pptx_builder, preview

SW, SH = 13.333, 7.5
MD_RE = re.compile(r"(^|\s)#{1,6}\s|\*\*")

STYLE = {"font_name": "微软雅黑", "title_size": 26, "body_size": 18,
         "title_color": "0066CC", "text_color": "2B3A4A", "accent": "0E5A8A",
         "cover_style": "centered", "cover": {"title_size": 40, "title_color": "0E5A8A"}}


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    """真实生成 10 页 deck（含 markdown 输入、卡片+要点、表格+图表、白字艺术字）。"""
    d = tmp_path_factory.mktemp("deck")
    deck_path = str(d / "deck.pptx")
    html_path = str(d / "deck.html")
    slides = [
        {"title": "封面", "bullets": ["全新智能解决方案"],
         "wordart": {"text": "智能未来 触手可及", "size": 44, "color": "0066CC"}},
        {"title": "一、产品概述",
         "cards": [{"title": "智能引擎", "desc": "搭载最新AI算法，实时学习用户习惯"},
                   {"title": "极速响应", "desc": "毫秒级响应速度，流畅无延迟"},
                   {"title": "安全可靠", "desc": "企业级数据加密，隐私全方位保护"}],
         "bullets": ["## 核心定位",
                     "新一代智能产品，融合AI技术与人性化设计，为用户带来前所未有的使用体验"]},
        {"title": "二、核心功能", "bullets": ["**智能识别**：AI图像识别准确率高达99.2%"]},
        {"title": "三、技术优势",
         "cards": [{"title": "自主研发", "desc": "核心技术自主可控"}],
         "bullets": ["自研深度学习框架，训练效率提升3倍"]},
        {"title": "四、市场表现",
         "table": {"header": ["指标", "2022", "2023", "2024"],
                   "rows": [["用户数(万)", 12, 28, 45]]},
         "chart": {"type": "column", "title": "用户数（万）",
                   "labels": ["2022", "2023", "2024"], "values": [12, 28, 45]},
         "bullets": ["## 行业地位", "连续3年保持细分市场占有率第一"]},
        {"title": "五、应用场景", "bullets": ["#### 电商零售", "转化率提升30%"]},
        {"title": "六、发展规划", "bullets": ["## 2025年目标", "- 用户规模突破100万"]},
        {"title": "七、合作共赢",
         "bullets": ["## 合作伙伴计划", "共建智能生态"],
         "wordart": {"text": "携手共赢 共创未来", "size": 36, "color": "0066CC"}},
        {"title": "结语", "bullets": ["期待与您合作", "共筑智能新时代"],
         "wordart": {"text": "THANK YOU", "size": 48, "color": "FFFFFF"}},
    ]
    pptx_builder.build_pptx(deck_path, "智能产品发布会", slides, STYLE)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(preview.render_office_html(deck_path) or "")
    return {"deck": deck_path, "html": html_path}


def _page_text(slide):
    out = []
    for sh in slide.shapes:
        try:
            if sh.has_text_frame and sh.has_text_frame:
                out.append(sh.text_frame.text)
        except Exception:
            pass
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                out.extend(c.text for c in r.cells)
    return "\n".join(out)


def _body_boxes(slide):
    """正文块（排除标题带与页码块），用于遮挡检查。"""
    boxes = []
    for sh in slide.shapes:
        if not (sh.has_text_frame and sh.has_text_frame):
            continue
        if not sh.text_frame.text.strip():
            continue
        x, y = Emu(sh.left or 0).inches, Emu(sh.top or 0).inches
        w, h = Emu(sh.width or 0).inches, Emu(sh.height or 0).inches
        if x > 11.5 or y < 1.3:
            continue
        boxes.append((sh.shape_id, x, y, w, h))
    return boxes


def test_no_markdown_residue(deck):
    """### / #### / ** 不得原样出现在成品文案里。"""
    prs = Presentation(deck["deck"])
    for si, slide in enumerate(prs.slides, 1):
        for line in _page_text(slide).splitlines():
            assert not MD_RE.search(line), f"第{si}页残留 markdown: {line!r}"


def test_markdown_levels_become_styled_runs(deck):
    """'## 核心定位' → 层级标题（加粗 + 主色）；'**x**' → 行内加粗 run。"""
    prs = Presentation(deck["deck"])
    bold_title = bold_inline = False
    for slide in prs.slides:
        for sh in slide.shapes:
            try:
                if not (sh.has_text_frame and sh.has_text_frame):
                    continue
            except Exception:
                continue
            for p in sh.text_frame.paragraphs:
                runs = list(p.runs)
                # 子标题：段落文本即标题本身且加粗
                if p.text.strip() == "核心定位" and any(r.font.bold for r in runs):
                    bold_title = True
                # 行内粗体：仅「智能识别」这一段是加粗 run（星号已被剥离）
                if any(r.font.bold and r.text.strip() == "智能识别" for r in runs):
                    bold_inline = True
    assert bold_title, "Markdown 子标题未转为加粗标题 run"
    assert bold_inline, "行内 **粗体** 未转为加粗 run"


def test_body_blocks_not_overlapping(deck):
    """卡片与要点、艺术字与要点不得互相压叠。"""
    prs = Presentation(deck["deck"])
    for si, slide in enumerate(prs.slides, 1):
        boxes = _body_boxes(slide)
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
                assert not (ox > 0.25 and oy > 0.25), \
                    f"第{si}页正文块重叠 id{a[0]}×id{b[0]} ({ox:.2f}×{oy:.2f}in)"


def test_text_inside_page(deck):
    """正文文本框不得越出页面。"""
    prs = Presentation(deck["deck"])
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if str(sh.shape_type or "").startswith("TEXT_BOX") is False:
                continue
            x, y = Emu(sh.left or 0).inches, Emu(sh.top or 0).inches
            w, h = Emu(sh.width or 0).inches, Emu(sh.height or 0).inches
            assert x >= -0.02 and y >= -0.02, f"第{si}页文本框越界左上 {x:.2f},{y:.2f}"
            assert x + w <= SW + 0.02 and y + h <= SH + 0.02, \
                f"第{si}页文本框越界右下 {x + w:.2f},{y + h:.2f}"


def test_invisible_text_autofixed(deck):
    """白底上的白字（#FFFFFF 艺术字）须被换成可读色。"""
    prs = Presentation(deck["deck"])
    white_art = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            try:
                if not (sh.has_text_frame and sh.has_text_frame):
                    continue
            except Exception:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() == "THANK YOU":
                        col = r.font.color
                        white_art.append(str(col.rgb) if col and col.type is not None else "")
    assert white_art, "未找到结语艺术字"
    assert all(c.upper() != "FFFFFF" for c in white_art), \
        f"白底白字未修复: {white_art}"


def test_table_and_chart_stack_on_same_slide(deck):
    """表格 + 柱状图同页时二者都要渲染，且纵向不重叠。"""
    prs = Presentation(deck["deck"])
    slide = list(prs.slides)[5]        # 封面 1 + 4 页 → 第 5 张内容页
    has_table = any(getattr(sh, "has_table", False) and sh.has_table for sh in slide.shapes)
    has_bar = any(sh.shape_type is not None and str(sh.shape_type).startswith("AUTO_SHAPE")
                  and Emu(sh.width or 0).inches < 1.2 for sh in slide.shapes)
    assert has_table, "表格未渲染"
    assert has_bar, "柱状图未渲染（表格+图表同页堆叠失效）"
    rects = [(Emu(sh.top or 0).inches, Emu(sh.height or 0).inches)
             for sh in slide.shapes
             if str(sh.shape_type or "").startswith("TABLE")]
    bars = [(Emu(sh.top or 0).inches, Emu(sh.height or 0).inches)
            for sh in slide.shapes
            if str(sh.shape_type or "").startswith("AUTO_SHAPE")
            and Emu(sh.width or 0).inches < 1.2 and Emu(sh.height or 0).inches > 0.1]
    for ty, th in rects:
        for by, bh in bars:
            overlap = min(ty + th, by + bh) - max(ty, by)
            assert overlap <= 0.25, f"表格与柱状图重叠 {overlap:.2f}in"


def test_page_number_block_geometry(deck):
    """页码块须为 1.00×0.40in（96.0×38.4px），文本为页码数字。"""
    prs = Presentation(deck["deck"])
    for si, slide in enumerate(prs.slides, 1):
        found = False
        for sh in slide.shapes:
            if not (sh.has_text_frame and sh.has_text_frame):
                continue
            if sh.text_frame.text.strip() != str(si):
                continue
            w, h = Emu(sh.width or 0).inches, Emu(sh.height or 0).inches
            assert abs(w - 1.00) < 0.01 and abs(h - 0.40) < 0.01, \
                f"第{si}页页码块尺寸 {w:.2f}×{h:.2f}in ≠ 1.00×0.40in"
            found = True
        assert found, f"第{si}页缺少页码块"


def test_html_spec_and_scripts(deck):
    with open(deck["html"], encoding="utf-8") as f:
        h = f.read()
    body = h.split("<script>")[0]
    n = h.count("<div class='deck' data-fit")
    assert n == 10, f"页数应为 10，实际 {n}"
    assert "class='sg'" in h and "class='tb'" in h, "缺少 .sg / .tb 图层"
    assert "width:1280px" in h and "height:720px" in h, "页面尺寸非 1280×720"
    assert "width:96.0px;height:38.4px" in h, "页码块未按 96.0×38.4px 输出"
    nums = re.findall(r"class='slidenum'>(\d+) / (\d+)</div>", h)
    assert [int(a) for a, _ in nums] == list(range(1, 11)), f"slidenum 序号异常 {nums}"
    assert all(b == "10" for _, b in nums), "slidenum 分母应为 10"
    for frag in ("var hostW = host.clientWidth || 0;", "__fitSlots", "__syncSlots",
                 "ResizeObserver", "__deckShowPage", "__offaReady", "offa-loading"):
        assert frag in h, f"缩放/放映脚本缺片段: {frag}"
    assert "###" not in body and "####" not in body and "**" not in body, \
        "HTML 正文残留 markdown 标记"


def test_fit_slots_keeps_only_current_page(deck):
    """多页画布只保留当前页占位，避免下方堆出大片空白（"预览空白"观感）。"""
    with open(deck["html"], encoding="utf-8") as f:
        js = f.read().split("<script>")[1].split("</script>")[0]
    assert "function __syncSlots" in js, "缺少槽位同步函数"
    assert "classList.contains('on') ? '' : 'none'" in js, "槽位未按当前页收起"
    # 翻页必须同步槽位，否则切页后新页也保持隐藏
    show = js.split("function __deckShowPage")[1].split("function ")[0]
    assert "__syncSlots()" in show, "翻页未同步槽位可见性"


# ---------------------------------------------------------------------------
# 图形排版：形状不得「部分压叠」（历史缺陷：循环图环上块压住中央块、
# 表格底板高度把 +padding 写在 min 之外导致侵入下一个块）
# ---------------------------------------------------------------------------
def _shape_rects(slide):
    out = []
    for sh in slide.shapes:
        x, y = Emu(sh.left or 0).inches, Emu(sh.top or 0).inches
        w, h = Emu(sh.width or 0).inches, Emu(sh.height or 0).inches
        if w * h <= 0.05 or w < 0.1 or h < 0.14:
            continue
        if x < -0.02 or y < -0.02 or x + w > SW + 0.02 or y + h > SH + 0.02:
            continue
        if x > 11.5 or y < 1.3:
            continue
        out.append((sh.shape_id, x, y, w, h))
    return out


def _contained(a, b, eps=0.04):
    return (a[1] >= b[1] - eps and a[2] >= b[2] - eps
            and a[1] + a[3] <= b[1] + b[3] + eps and a[2] + a[4] <= b[2] + b[4] + eps)


@pytest.mark.parametrize("diagram", [
    {"type": "cycle", "items": ["技术创新", "产品迭代", "市场拓展", "用户服务"]},
    {"type": "cycle", "center": "闭环", "items": ["A", "B", "C", "D", "E"]},
    {"type": "flow", "items": ["数据采集", "智能分析", "结果呈现", "持续优化"]},
])
def test_diagram_shapes_not_partially_overlapping(tmp_path, diagram):
    out = str(tmp_path / "d.pptx")
    pptx_builder.build_pptx(out, "图形排版", [
        {"title": "发展规划", "diagram": diagram,
         "bullets": ["## 目标", "规模突破100万", "进入3个新细分市场"]},
    ], STYLE)
    prs = Presentation(out)
    slide = prs.slides[1]
    rects = _shape_rects(slide)
    bad = []
    for i in range(len(rects)):
        for j in range(i + 1, len(rects)):
            a, b = rects[i], rects[j]
            ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
            if ox > 0.08 and oy > 0.08 and not (_contained(a, b) or _contained(b, a)):
                bad.append((a[0], b[0], round(ox, 2), round(oy, 2)))
    assert not bad, f"{diagram.get('type')} 形状部分压叠: {bad[:3]}"


def test_table_background_stays_in_budget(tmp_path):
    """表格底板不得越过本块高度预算（否则会压到同页下一个块的标题）。"""
    out = str(tmp_path / "t.pptx")
    pptx_builder.build_pptx(out, "", [
        {"title": "市场表现",
         "table": {"header": ["指标", "2022", "2023", "2024"],
                   "rows": [["营收(亿)", 1.2, 2.8, 4.5], ["用户数(万)", 12, 28, 45],
                            ["NPS", 72, 81, 87]]},
         "chart": {"type": "column", "title": "用户数（万）",
                   "labels": ["2022", "2023", "2024"], "values": [12, 28, 45]},
         "bullets": ["## 行业地位", "占有率第一"]},
    ], STYLE)
    prs = Presentation(out)
    slide = prs.slides[0]
    chart_title = None
    for sh in slide.shapes:
        try:
            if sh.has_text_frame and sh.text_frame.text.strip() == "用户数（万）":
                chart_title = (Emu(sh.top or 0).inches, Emu(sh.left or 0).inches,
                               Emu(sh.width or 0).inches)
        except Exception:
            continue
    assert chart_title, "未找到图表标题"
    ty, tx, tw = chart_title
    for sh in slide.shapes:
        if str(sh.shape_type or "").startswith("AUTO_SHAPE") is False:
            continue
        x, y = Emu(sh.left or 0).inches, Emu(sh.top or 0).inches
        w, h = Emu(sh.width or 0).inches, Emu(sh.height or 0).inches
        if abs(w - tw) > 0.5 or x > tx + 0.5:      # 只看与表格同宽的底板
            continue
        if h < 0.14:                              # 排除图表坐标轴/网格细线
            continue
        assert y + h <= ty + 0.02, \
            f"底板 {sh.shape_id} 底部 {y + h:.2f} 越过图表标题顶部 {ty:.2f}"
